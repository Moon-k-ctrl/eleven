"""Word (.docx) preview widget using python-docx."""
from __future__ import annotations

import logging
from pathlib import Path

from PyQt6.QtWidgets import QTextBrowser, QVBoxLayout, QWidget

logger = logging.getLogger("eleven.preview.word")


class WordPreview(QWidget):
    """DOCX preview rendered as HTML in QTextBrowser."""

    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)

        self._browser = QTextBrowser()
        self._browser.setOpenExternalLinks(False)
        self._browser.setStyleSheet("""
            QTextBrowser {
                background: #1e1e1e;
                color: #e0e0e0;
                border: none;
                padding: 16px;
                font-size: 14px;
            }
        """)
        layout.addWidget(self._browser)

    def load_docx(self, path: str) -> None:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document(path)
            html_parts = [
                "<style>"
                "body { background: #1e1e1e; color: #e0e0e0; font-family: 'Microsoft YaHei', sans-serif; }"
                "h1, h2, h3 { color: #fff; }"
                "table { border-collapse: collapse; width: 100%; margin: 8px 0; }"
                "th, td { border: 1px solid #444; padding: 6px 10px; text-align: left; }"
                "th { background: #333; }"
                "p { margin: 4px 0; line-height: 1.6; }"
                "code { background: #2a2a2a; padding: 2px 4px; border-radius: 3px; }"
                "</style>"
            ]

            for para in doc.paragraphs:
                style_name = para.style.name if para.style else ""
                text = self._runs_to_html(para)
                if not text.strip():
                    html_parts.append("<br>")
                elif style_name.startswith("Heading 1"):
                    html_parts.append(f"<h1>{text}</h1>")
                elif style_name.startswith("Heading 2"):
                    html_parts.append(f"<h2>{text}</h2>")
                elif style_name.startswith("Heading 3"):
                    html_parts.append(f"<h3>{text}</h3>")
                elif style_name.startswith("List"):
                    html_parts.append(f"<li>{text}</li>")
                else:
                    html_parts.append(f"<p>{text}</p>")

            for table in doc.tables:
                html_parts.append("<table>")
                for i, row in enumerate(table.rows):
                    tag = "th" if i == 0 else "td"
                    cells = "".join(
                        f"<{tag}>{cell.text}</{tag}>" for cell in row.cells
                    )
                    html_parts.append(f"<tr>{cells}</tr>")
                html_parts.append("</table>")

            self._browser.setHtml("\n".join(html_parts))

        except ImportError:
            self._browser.setHtml(
                '<p style="color: #ff6b6b;">需要安装 python-docx: pip install python-docx</p>'
            )
        except Exception as e:
            logger.error(f"Failed to load DOCX: {e}")
            self._browser.setHtml(f'<p style="color: #ff6b6b;">加载失败: {e}</p>')

    @staticmethod
    def _runs_to_html(para) -> str:
        parts = []
        for run in para.runs:
            text = run.text
            if not text:
                continue
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if run.bold:
                text = f"<b>{text}</b>"
            if run.italic:
                text = f"<i>{text}</i>"
            if run.underline:
                text = f"<u>{text}</u>"
            parts.append(text)
        return "".join(parts)
