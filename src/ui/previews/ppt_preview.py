"""PowerPoint (.ppt/.pptx) preview using PowerPoint COM for 100% fidelity."""
from __future__ import annotations

import logging
import shutil
import tempfile
from pathlib import Path
from typing import Optional

from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QPixmap
from PyQt6.QtWidgets import (
    QLabel,
    QScrollArea,
    QVBoxLayout,
    QWidget,
)

logger = logging.getLogger("eleven.preview.ppt")


class PptPreview(QWidget):
    """PowerPoint slide viewer — COM-rendered slides as images."""

    slide_changed = pyqtSignal(int, int)  # current, total

    def __init__(self, parent=None):
        super().__init__(parent)
        self._slide_pixmaps: list[QPixmap] = []
        self._current_slide = 0
        self._total_slides = 0
        self._export_dir: Optional[Path] = None

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        self._image_label = QLabel()
        self._image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._image_label.setStyleSheet("background: #1e1e1e;")

        self._scroll = QScrollArea()
        self._scroll.setWidget(self._image_label)
        self._scroll.setWidgetResizable(True)
        self._scroll.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self._scroll, stretch=1)

    def load_pptx(self, path: str) -> None:
        """Load PPT/PPTX via PowerPoint COM, exporting each slide as PNG."""
        self._export_slides_com(path)
        if self._slide_pixmaps:
            self._show_slide(0)
        self.slide_changed.emit(
            self._current_slide + 1 if self._slide_pixmaps else 0,
            self._total_slides,
        )

    def _export_slides_com(self, path: str) -> None:
        """Use PowerPoint COM to export slides as PNG images."""
        ppt_app = None
        presentation = None
        try:
            import win32com.client

            ppt_app = win32com.client.Dispatch("PowerPoint.Application")
            presentation = ppt_app.Presentations.Open(
                str(Path(path).resolve()),
                ReadOnly=True,
                Untitled=True,
                WithWindow=False,
            )

            self._total_slides = presentation.Slides.Count
            self._export_dir = Path(tempfile.mkdtemp(prefix="ppt_preview_"))

            for i in range(1, self._total_slides + 1):
                slide = presentation.Slides(i)
                img_path = self._export_dir / f"slide_{i:03d}.png"
                slide.Export(str(img_path), "PNG", 1920, 1080)
                pixmap = QPixmap(str(img_path))
                if not pixmap.isNull():
                    self._slide_pixmaps.append(pixmap)
                else:
                    logger.warning(f"Slide {i}: export produced null pixmap")

        except ImportError:
            logger.error("win32com not available, falling back to python-pptx")
            self._fallback_text_extract(path)
        except Exception as e:
            logger.error(f"COM export failed: {e}", exc_info=True)
            self._fallback_text_extract(path)
        finally:
            if presentation:
                try:
                    presentation.Close()
                except Exception:
                    pass
            if ppt_app:
                try:
                    ppt_app.Quit()
                except Exception:
                    pass

    def _fallback_text_extract(self, path: str) -> None:
        """COM failure fallback: extract text + images via python-pptx."""
        try:
            from pptx import Presentation
            from pptx.util import Emu
            from PyQt6.QtGui import QImage

            prs = Presentation(path)
            self._total_slides = len(prs.slides)

            # Slide dimensions (EMU)
            slide_w = prs.slide_width
            slide_h = prs.slide_height

            for slide in prs.slides:
                pixmap = QPixmap(1920, 1080)
                pixmap.fill(Qt.GlobalColor.white)
                painter = QPainter(pixmap)

                scale_x = 1920.0 / slide_w
                scale_y = 1080.0 / slide_h

                for shape in slide.shapes:
                    sx = int(shape.left * scale_x) if shape.left else 0
                    sy = int(shape.top * scale_y) if shape.top else 0
                    sw = int(shape.width * scale_x) if shape.width else 200
                    sh = int(shape.height * scale_y) if shape.height else 50

                    # Extract images
                    if shape.shape_type == 13:  # PICTURE
                        try:
                            image = shape.image
                            qimage = QImage.fromData(image.blob)
                            if not qimage.isNull():
                                img_pix = QPixmap.fromImage(qimage)
                                scaled = img_pix.scaled(
                                    sw, sh,
                                    Qt.AspectRatioMode.KeepAspectRatio,
                                    Qt.TransformationMode.SmoothTransformation,
                                )
                                painter.drawPixmap(sx, sy, scaled)
                        except Exception:
                            pass

                    # Extract text
                    elif shape.has_text_frame:
                        from PyQt6.QtGui import QFont
                        font = QFont("Microsoft YaHei", 14)
                        painter.setFont(font)
                        painter.setPen(Qt.GlobalColor.black)
                        y = sy + 20
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                # Truncate long lines
                                display = t[:80] + ("..." if len(t) > 80 else "")
                                painter.drawText(sx + 4, y, display)
                                y += 24
                                if y > 1060:
                                    break

                painter.end()
                self._slide_pixmaps.append(pixmap)

        except ImportError:
            logger.error("python-pptx not available for fallback")
        except Exception as e:
            logger.error(f"Fallback text extract failed: {e}")

    def _show_slide(self, index: int) -> None:
        if not (0 <= index < len(self._slide_pixmaps)):
            return
        self._current_slide = index
        pixmap = self._slide_pixmaps[index]
        available = self._scroll.viewport().size()
        scaled = pixmap.scaled(
            available,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        self._image_label.setPixmap(scaled)
        self.slide_changed.emit(index + 1, self._total_slides)

    def next_slide(self) -> None:
        if self._current_slide < self._total_slides - 1:
            self._show_slide(self._current_slide + 1)

    def prev_slide(self) -> None:
        if self._current_slide > 0:
            self._show_slide(self._current_slide - 1)

    def goto_slide(self, index: int) -> None:
        if 0 <= index < self._total_slides:
            self._show_slide(index)

    @property
    def total_slides(self) -> int:
        return self._total_slides

    def resizeEvent(self, event) -> None:
        super().resizeEvent(event)
        if self._slide_pixmaps:
            self._show_slide(self._current_slide)

    def cleanup(self) -> None:
        """Remove temporary export directory."""
        if self._export_dir and self._export_dir.exists():
            try:
                shutil.rmtree(self._export_dir, ignore_errors=True)
            except Exception:
                pass
