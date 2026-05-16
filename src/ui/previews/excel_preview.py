"""Excel (.xlsx) preview widget using openpyxl."""
from __future__ import annotations

import logging
from pathlib import Path

from PyQt6.QtWidgets import QTextBrowser, QVBoxLayout, QWidget

logger = logging.getLogger("eleven.preview.excel")

# Max rows to render (avoid huge sheets freezing the UI)
_MAX_ROWS = 500


class ExcelPreview(QWidget):
    """XLSX preview rendered as an HTML table in QTextBrowser."""

    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)

        self._browser = QTextBrowser()
        self._browser.setOpenExternalLinks(False)
        self._browser.setStyleSheet("""
            QTextBrowser {
                background: #1e1e1e;
                color: #e0e0e0;
                border: none;
                padding: 16px;
                font-size: 13px;
            }
        """)
        layout.addWidget(self._browser)

    def load_xlsx(self, path: str) -> None:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            sheet = wb.active
            if sheet is None:
                self._browser.setHtml(
                    '<p style="color: #aaa;">工作簿中没有工作表</p>'
                )
                return

            html_parts = [
                "<style>"
                "body { background: #1e1e1e; color: #e0e0e0; "
                "font-family: 'Microsoft YaHei', sans-serif; }"
                "table { border-collapse: collapse; width: 100%; margin: 8px 0; }"
                "th, td { border: 1px solid #444; padding: 4px 8px; text-align: left; "
                "white-space: nowrap; max-width: 300px; overflow: hidden; "
                "text-overflow: ellipsis; }"
                "th { background: #333; color: #fff; font-weight: bold; }"
                "tr:nth-child(even) { background: #252525; }"
                "tr:hover { background: #2a2a2a; }"
                ".sheet-name { color: #888; font-size: 12px; margin-bottom: 8px; }"
                ".row-count { color: #666; font-size: 11px; margin-top: 8px; }"
                "</style>",
                f'<div class="sheet-name">Sheet: {sheet.title}</div>',
                "<table>",
            ]

            row_count = 0
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= _MAX_ROWS:
                    break
                tag = "th" if i == 0 else "td"
                cells = []
                for val in row:
                    text = "" if val is None else str(val)
                    text = (text.replace("&", "&amp;")
                                .replace("<", "&lt;")
                                .replace(">", "&gt;"))
                    cells.append(f"<{tag}>{text}</{tag}>")
                html_parts.append(f"<tr>{''.join(cells)}</tr>")
                row_count += 1

            html_parts.append("</table>")

            total_rows = sheet.max_row or 0
            if total_rows > _MAX_ROWS:
                html_parts.append(
                    f'<div class="row-count">显示前 {_MAX_ROWS} 行，'
                    f'共 {total_rows} 行</div>'
                )
            else:
                html_parts.append(f'<div class="row-count">共 {row_count} 行</div>')

            self._browser.setHtml("\n".join(html_parts))
            wb.close()

        except ImportError:
            self._browser.setHtml(
                '<p style="color: #ff6b6b;">需要安装 openpyxl: '
                'pip install openpyxl</p>'
            )
        except Exception as e:
            logger.error(f"Failed to load XLSX: {e}")
            self._browser.setHtml(
                f'<p style="color: #ff6b6b;">加载失败: {e}</p>'
            )
